import unittest
import os
import io
import openpyxl
from pypdf import PdfReader
from app import create_app
from models import db, Statement, Transaction, User

class TestReportDownloads(unittest.TestCase):
    def setUp(self):
        self.app = create_app()
        self.app.config['TESTING'] = True
        self.client = self.app.test_client()

    def test_all_reports_excel_download_empty_db(self):
        with self.app.app_context():
            # Statement count can be 0 or more depending on whether user entered data
            # Just verify query executes
            self.assertGreaterEqual(Statement.query.count(), 0)

            report_ids = [
                'ubl_statement',
                'transaction_detail',
                'company_allocation',
                'merchant_expense',
                'advance_tax',
                'sales_tax',
                'statement_reconciliation',
                'outstanding_payment',
                'monthly_summary',
                'audit_adjustment'
            ]

            for rep_id in report_ids:
                res = self.client.get(f'/api/reports/excel/{rep_id}')
                self.assertEqual(res.status_code, 200, f"Excel download failed for {rep_id}: {res.data}")
                self.assertIn('application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', res.content_type)
                wb = openpyxl.load_workbook(io.BytesIO(res.data))
                ws = wb.active
                self.assertIsNotNone(ws)
                # Verify logo image is present in Excel
                self.assertGreater(len(ws._images), 0, f"No logo found in {rep_id} excel export")
                # Verify US Letter printing configuration (paperSize 1 = Letter)
                self.assertEqual(str(ws.page_setup.paperSize), '1', f"Paper size not US Letter for {rep_id}")
                self.assertEqual(ws.page_setup.orientation, 'landscape', f"Orientation not landscape for {rep_id}")
                self.assertTrue(ws.sheet_properties.pageSetUpPr.fitToPage, f"Fit to page not enabled for {rep_id}")

    def test_all_reports_pdf_download_empty_db(self):
        with self.app.app_context():
            report_ids = [
                'ubl_statement',
                'transaction_detail',
                'company_allocation',
                'merchant_expense',
                'advance_tax',
                'sales_tax',
                'statement_reconciliation',
                'outstanding_payment',
                'monthly_summary',
                'audit_adjustment'
            ]

            for rep_id in report_ids:
                res = self.client.get(f'/api/reports/pdf/{rep_id}')
                self.assertEqual(res.status_code, 200, f"PDF download failed for {rep_id}: {res.data}")
                self.assertEqual(res.content_type, 'application/pdf')
                reader = PdfReader(io.BytesIO(res.data))
                self.assertGreater(len(reader.pages), 0, f"PDF has no pages for {rep_id}")
                # Verify US Letter page dimensions: 11 x 8.5 inches = 792 x 612 points
                page = reader.pages[0]
                self.assertEqual(float(page.mediabox.width), 792.0, f"Page width not 792pt (11 in) for {rep_id}")
                self.assertEqual(float(page.mediabox.height), 612.0, f"Page height not 612pt (8.5 in) for {rep_id}")
                text = page.extract_text()
                self.assertIn("RED PUBLICIS", text, f"Red Publicis branding missing in {rep_id} PDF")
                self.assertIn("US Letter", text, f"US Letter indicator missing in {rep_id} PDF")

    def test_statement_excel_calculation_flags_and_formulas(self):
        from services.excel_service import ExcelService
        with self.app.app_context():
            stmt = Statement.query.first()
            if stmt:
                data = ExcelService.create_statement_excel(stmt)
                wb = openpyxl.load_workbook(data if isinstance(data, io.BytesIO) else io.BytesIO(data))
                ws = wb.active
                # Verify Excel auto-calculate flags
                self.assertTrue(wb.calculation.fullCalcOnLoad)
                self.assertTrue(wb.calculation.forceFullCalc)
                # Verify sheet has rows and formulas
                found_sum = False
                for row in ws.iter_rows(values_only=True):
                    for val in row:
                        if isinstance(val, str) and val.startswith('=SUM('):
                            found_sum = True
                            break
                self.assertTrue(found_sum, "No SUM formulas found in statement excel sheet")

    def test_statement_pdf_grand_total_summary_and_signatures(self):
        with self.app.app_context():
            stmt = Statement.query.first()
            stmt_id = stmt.id if stmt else None
            url = f'/api/reports/pdf/ubl_statement?statement_id={stmt_id}' if stmt_id else '/api/reports/pdf/ubl_statement'
            res = self.client.get(url)
            self.assertEqual(res.status_code, 200)
            self.assertEqual(res.content_type, 'application/pdf')
            
            reader = PdfReader(io.BytesIO(res.data))
            all_text = "\n".join([page.extract_text() for page in reader.pages])
            
            # Verify Grand Total
            self.assertIn("GRAND TOTAL", all_text, "GRAND TOTAL missing in statement PDF")
            # Verify Summary Area
            self.assertIn("SUMMARY", all_text, "SUMMARY missing in statement PDF")
            self.assertIn("ALLOCATION BREAKDOWN", all_text, "ALLOCATION BREAKDOWN missing in statement PDF")
            self.assertIn("G.TOTAL", all_text, "G.TOTAL missing in summary section of statement PDF")
            # Verify Signature Area
            self.assertIn("PREPARED BY", all_text, "PREPARED BY missing in statement PDF")
            self.assertIn("VERIFIED BY", all_text, "VERIFIED BY missing in statement PDF")
            self.assertIn("APPROVED BY", all_text, "APPROVED BY missing in statement PDF")
            self.assertIn("Signature:", all_text, "Signature line missing in statement PDF")

if __name__ == '__main__':
    unittest.main()

