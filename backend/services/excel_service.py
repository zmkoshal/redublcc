import os
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.units import pixels_to_EMU
from decimal import Decimal
from datetime import datetime

class ExcelService:
    @staticmethod
    def get_logo_path():
        candidates = [
            os.path.join(os.path.dirname(__file__), '..', 'static', 'logo.png'),
            os.path.join(os.path.dirname(__file__), '..', '..', 'frontend', 'public', 'logo.png'),
            os.path.join(os.path.dirname(__file__), '..', '..', 'frontend', 'dist', 'logo.png'),
        ]
        for p in candidates:
            if os.path.exists(p):
                return os.path.abspath(p)
        return None

    @staticmethod
    def create_statement_excel(statement=None, transactions=None, allocation_categories=None):
        """
        Creates an exact Excel workbook matching the structure, grouping,
        subtotals, formulas, grand totals, and summary of the original UBL Credit Card sheet.
        """
        if statement is not None and transactions is None:
            transactions = getattr(statement, 'transactions', [])
        if transactions is None:
            transactions = []
        if allocation_categories is None:
            try:
                from models import AllocationCategory
                allocation_categories = AllocationCategory.query.order_by(AllocationCategory.display_order.asc()).all()
            except Exception:
                allocation_categories = []

        wb = openpyxl.Workbook()
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        ws = wb.active
        ws.title = "Statement"
        ws.views.sheetView[0].showGridLines = True

        # Styles
        title_font = Font(name="Calibri", size=14, bold=True, color="002060")
        subtitle_font = Font(name="Calibri", size=11, bold=True, color="1F4E79")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        
        group_header_font = Font(name="Calibri", size=11, bold=True, color="002060")
        group_header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        
        subtotal_font = Font(name="Calibri", size=11, bold=True)
        subtotal_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
        
        grand_total_font = Font(name="Calibri", size=12, bold=True, color="000000")
        grand_total_fill = PatternFill(start_color="C6D9F1", end_color="C6D9F1", fill_type="solid")
        
        summary_header_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")

        thin_side = Side(border_style="thin", color="D9D9D9")
        thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        double_bottom = Border(top=thin_side, bottom=Side(border_style="double", color="000000"))

        # Row heights for heading area
        ws.row_dimensions[1].height = 26
        ws.row_dimensions[2].height = 22
        ws.row_dimensions[3].height = 10
        ws.row_dimensions[4].height = 28

        # Fallback if statement is None
        if statement is None:
            class DummyStatement:
                statement_number = "TEMPLATE"
                statement_date = datetime.now().date()
                payment_due_date = datetime.now().date()
                total_statement_amount = Decimal('0.00')
                amount_paid = Decimal('0.00')
                outstanding_balance = Decimal('0.00')
                status = "Draft"
            statement = DummyStatement()

        # Row 1: Title (aligned on left, with logo on right)
        ws.cell(row=1, column=2, value="RED PUBLICIS - UBL CREDIT CARD STATEMENT FOR PAYMENT").font = title_font
        
        # Row 2: Statement Date and Due Date
        stmt_date_str = statement.statement_date.strftime('%d-%b-%Y') if statement.statement_date else datetime.now().strftime('%d-%b-%Y')
        due_date_str = statement.payment_due_date.strftime('%d-%b-%Y') if statement.payment_due_date else datetime.now().strftime('%d-%b-%Y')
        ws.cell(row=2, column=2, value=f"Statement Date: {stmt_date_str}   Payment Due Date: {due_date_str}").font = subtitle_font

        # Logo on right side: aligned directly with the heading, flush to right edge of Column I (the table's right border)
        logo_path = ExcelService.get_logo_path()
        if logo_path:
            try:
                img = OpenpyxlImage(logo_path)
                img.width = 170
                img.height = 34
                size = XDRPositiveSize2D(pixels_to_EMU(img.width), pixels_to_EMU(img.height))
                # Anchor in Column H (col=7) with 78px offset so right edge aligns flush with right border of Column I
                marker = AnchorMarker(col=7, colOff=pixels_to_EMU(78), row=0, rowOff=pixels_to_EMU(6))
                img.anchor = OneCellAnchor(_from=marker, ext=size)
                ws.add_image(img)
            except Exception as e:
                print(f"Warning: could not add logo to Excel: {e}")

        # Row 4: Headers
        headers = [
            (2, "Party / Merchant Name"),
            (3, "Online (USD)"),
            (4, "Exchange Rate"),
            (5, "Rs. (PKR)"),
            (6, "Adv Tax"),
            (7, "Sales Tax Withheld"),
            (8, "Total Amount"),
            (9, "Status")
        ]
        for col_idx, text in headers:
            cell = ws.cell(row=4, column=col_idx, value=text)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        current_row = 6
        subtotal_rows = []
        subtotal_map = {}

        # Group transactions by Allocation
        # Map allocation IDs to list of transactions
        groups = {}
        for cat in allocation_categories:
            groups[cat.id] = {'cat': cat, 'items': []}
        groups[None] = {'cat': None, 'items': []}

        for t in transactions:
            if t.allocation_id in groups:
                groups[t.allocation_id]['items'].append(t)
            else:
                groups[None]['items'].append(t)

        for cat_id, data in groups.items():
            items = data['items']
            if not items:
                continue

            cat_name = data['cat'].name if data['cat'] else "CHARGE TO UNALLOCATED"
            
            # Group Header Row
            g_cell = ws.cell(row=current_row, column=2, value=cat_name)
            g_cell.font = group_header_font
            g_cell.fill = group_header_fill
            for c in range(3, 10):
                ws.cell(row=current_row, column=c).fill = group_header_fill
            current_row += 1

            start_item_row = current_row

            # Items
            for t in items:
                ws.cell(row=current_row, column=2, value=t.description or t.merchant_name).border = thin_border
                
                c_usd = ws.cell(row=current_row, column=3, value=float(t.usd_amount) if t.usd_amount else None)
                if t.usd_amount:
                    c_usd.number_format = '#,##0.00'
                c_usd.border = thin_border

                if t.usd_amount:
                    c_rate = ws.cell(row=current_row, column=4, value=f"=E{current_row}/C{current_row}")
                else:
                    c_rate = ws.cell(row=current_row, column=4, value=float(t.exchange_rate) if t.exchange_rate else None)
                c_rate.number_format = '#,##0.0000'
                c_rate.border = thin_border

                c_pkr = ws.cell(row=current_row, column=5, value=float(t.pkr_amount or 0.0))
                c_pkr.number_format = '#,##0.00'
                c_pkr.border = thin_border

                c_adv = ws.cell(row=current_row, column=6, value=float(t.advance_tax or 0.0))
                c_adv.number_format = '#,##0.00'
                c_adv.border = thin_border

                c_st = ws.cell(row=current_row, column=7, value=float(t.sales_tax_withheld or 0.0))
                c_st.number_format = '#,##0.00'
                c_st.border = thin_border

                # Excel formula for Total: E + F + G
                c_tot = ws.cell(row=current_row, column=8, value=f"=SUM(E{current_row}:G{current_row})")
                c_tot.number_format = '#,##0.00'
                c_tot.font = Font(bold=True)
                c_tot.border = thin_border

                ws.cell(row=current_row, column=9, value=t.approval_status).border = thin_border
                current_row += 1

            end_item_row = current_row - 1

            # Subtotal Row
            subtotal_label = ws.cell(row=current_row, column=2, value=f"SUBTOTAL ({cat_name})")
            subtotal_label.font = subtotal_font
            subtotal_label.fill = subtotal_fill

            for c in range(3, 5):
                ws.cell(row=current_row, column=c).fill = subtotal_fill

            sub_pkr = ws.cell(row=current_row, column=5, value=f"=SUM(E{start_item_row}:E{end_item_row})")
            sub_pkr.number_format = '#,##0.00'
            sub_pkr.font = subtotal_font
            sub_pkr.fill = subtotal_fill

            sub_adv = ws.cell(row=current_row, column=6, value=f"=SUM(F{start_item_row}:F{end_item_row})")
            sub_adv.number_format = '#,##0.00'
            sub_adv.font = subtotal_font
            sub_adv.fill = subtotal_fill

            sub_st = ws.cell(row=current_row, column=7, value=f"=SUM(G{start_item_row}:G{end_item_row})")
            sub_st.number_format = '#,##0.00'
            sub_st.font = subtotal_font
            sub_st.fill = subtotal_fill

            sub_tot = ws.cell(row=current_row, column=8, value=f"=SUM(H{start_item_row}:H{end_item_row})")
            sub_tot.number_format = '#,##0.00'
            sub_tot.font = subtotal_font
            sub_tot.fill = subtotal_fill

            ws.cell(row=current_row, column=9).fill = subtotal_fill
            subtotal_rows.append(current_row)
            if cat_id is not None:
                subtotal_map[cat_id] = current_row
            current_row += 2  # spacing

        if not transactions:
            ws.cell(row=current_row, column=2, value="No transactions recorded yet in this statement").font = Font(name="Calibri", size=10, italic=True, color="595959")
            current_row += 2

        # Grand Total Row
        gt_row = current_row
        ws.cell(row=gt_row, column=2, value="GRAND TOTAL").font = grand_total_font
        ws.cell(row=gt_row, column=2).fill = grand_total_fill

        if subtotal_rows:
            pkr_sum_expr = f"SUM({', '.join([f'E{r}' for r in subtotal_rows])})"
            adv_sum_expr = f"SUM({', '.join([f'F{r}' for r in subtotal_rows])})"
            st_sum_expr = f"SUM({', '.join([f'G{r}' for r in subtotal_rows])})"
            tot_sum_expr = f"SUM({', '.join([f'H{r}' for r in subtotal_rows])})"
        else:
            pkr_sum_expr = "0"
            adv_sum_expr = "0"
            st_sum_expr = "0"
            tot_sum_expr = "0"

        for col_idx, expr in [(5, pkr_sum_expr), (6, adv_sum_expr), (7, st_sum_expr), (8, tot_sum_expr)]:
            c = ws.cell(row=gt_row, column=col_idx, value=f"={expr}")
            c.number_format = '#,##0.00'
            c.font = grand_total_font
            c.fill = grand_total_fill
            c.border = double_bottom

        current_row += 1

        # Payment Row
        pmt_row = current_row
        ws.cell(row=pmt_row, column=2, value="AMOUNT PAID").font = subtotal_font
        c_pmt = ws.cell(row=pmt_row, column=8, value=float(statement.amount_paid or 0.0))
        c_pmt.number_format = '#,##0.00'
        c_pmt.font = subtotal_font
        current_row += 1

        # Balance Due Row
        bal_row = current_row
        ws.cell(row=bal_row, column=2, value="OUTSTANDING BALANCE").font = grand_total_font
        c_bal = ws.cell(row=bal_row, column=8, value=f"=H{gt_row}-H{pmt_row}")
        c_bal.number_format = '#,##0.00'
        c_bal.font = grand_total_font
        c_bal.border = double_bottom
        current_row += 3

        # Signatures
        sig_row = current_row
        ws.cell(row=sig_row, column=2, value="IT / Prepared By: ___________________").font = Font(name="Calibri", size=10, bold=True)
        ws.cell(row=sig_row, column=5, value="Admin & HR: ___________________").font = Font(name="Calibri", size=10, bold=True)
        ws.cell(row=sig_row, column=8, value="Head of Finance: ___________________").font = Font(name="Calibri", size=10, bold=True)
        current_row += 3

        # SUMMARY SECTION (Mirroring rows 84-90 of original workbook)
        # Header banner spanning only from DETAIL (Col C) to AMOUNT (PKR) (Col E)
        ws.row_dimensions[current_row].height = 24
        for c in range(3, 6):
            ws.cell(row=current_row, column=c).fill = header_fill
        sum_title = ws.cell(row=current_row, column=3, value="SUMMARY")
        sum_title.font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
        sum_title.alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells(start_row=current_row, start_column=3, end_row=current_row, end_column=5)
        current_row += 1

        # Summary Column Headers: DETAIL (Cols C-D merged) and AMOUNT (PKR) (Col E)
        ws.row_dimensions[current_row].height = 22
        ws.merge_cells(start_row=current_row, start_column=3, end_row=current_row, end_column=4)
        h_det = ws.cell(row=current_row, column=3, value="DETAIL")
        h_det.font = Font(name="Calibri", size=11, bold=True, color="1F4E79")
        h_det.alignment = Alignment(horizontal="left", vertical="center")

        h_amt = ws.cell(row=current_row, column=5, value="AMOUNT (PKR)")
        h_amt.font = Font(name="Calibri", size=11, bold=True, color="1F4E79")
        h_amt.alignment = Alignment(horizontal="right", vertical="center")
        current_row += 1

        summary_start_row = current_row
        # Allocation summary rows dynamically referencing category subtotals
        for cat in allocation_categories:
            ws.row_dimensions[current_row].height = 20
            ws.merge_cells(start_row=current_row, start_column=3, end_row=current_row, end_column=4)
            c_name = ws.cell(row=current_row, column=3, value=cat.name)
            c_name.font = Font(name="Calibri", size=10, bold=True)
            c_name.alignment = Alignment(horizontal="left", vertical="center")

            if cat.id in subtotal_map:
                c_val = ws.cell(row=current_row, column=5, value=f"=H{subtotal_map[cat.id]}")
            else:
                c_val = ws.cell(row=current_row, column=5, value=0.0)
            c_val.number_format = '#,##0.00'
            c_val.font = Font(name="Calibri", size=10, bold=True)
            c_val.alignment = Alignment(horizontal="right", vertical="center")
            current_row += 1

        summary_end_row = current_row - 1
        # G.TOTAL of Summary
        ws.row_dimensions[current_row].height = 22
        c_lbl = ws.cell(row=current_row, column=4, value="G.TOTAL")
        c_lbl.font = Font(name="Calibri", size=11, bold=True)
        c_lbl.alignment = Alignment(horizontal="right", vertical="center")

        if summary_end_row >= summary_start_row:
            sum_tot = ws.cell(row=current_row, column=5, value=f"=SUM(E{summary_start_row}:E{summary_end_row})")
        else:
            sum_tot = ws.cell(row=current_row, column=5, value=0.0)
        sum_tot.font = Font(name="Calibri", size=11, bold=True)
        sum_tot.number_format = '#,##0.00'
        sum_tot.alignment = Alignment(horizontal="right", vertical="center")
        sum_tot.border = double_bottom

        # Column widths adjustment
        column_widths = {
            'A': 4,
            'B': 45,
            'C': 16,
            'D': 16,
            'E': 18,
            'F': 14,
            'G': 18,
            'H': 20,
            'I': 14
        }
        for col, width in column_widths.items():
            ws.column_dimensions[col].width = width

        # Print Setup for US Letter (8.5 x 11 in) Landscape Printing
        ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
        ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_margins.left = 0.4
        ws.page_margins.right = 0.4
        ws.page_margins.top = 0.5
        ws.page_margins.bottom = 0.5
        ws.print_options.horizontalCentered = True
        ws.print_title_rows = '4:4'
        ws.oddFooter.left.text = "RED PUBLICIS — UBL Statement (US Letter 8.5 x 11 in)"
        ws.oddFooter.right.text = "Page &P of &N"

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def create_generic_report(report_title, headers, rows, filename):
        """Generates clean, styled Excel export for any of the 10 financial reports."""
        wb = openpyxl.Workbook()
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        ws = wb.active
        ws.title = "Report"
        ws.views.sheetView[0].showGridLines = True

        title_font = Font(name="Calibri", size=14, bold=True, color="002060")

        ws.row_dimensions[1].height = 24
        ws.row_dimensions[2].height = 20
        ws.row_dimensions[3].height = 16
        ws.row_dimensions[4].height = 10

        ws.cell(row=1, column=1, value="RED PUBLICIS (COMMUNICATION ARTS)").font = title_font
        ws.cell(row=2, column=1, value=report_title).font = Font(name="Calibri", size=12, bold=True, color="1F4E79")
        ws.cell(row=3, column=1, value=f"Generated on: {datetime.now().strftime('%d-%b-%Y %H:%M')}").font = Font(size=9, color="595959")

        logo_path = ExcelService.get_logo_path()
        if logo_path:
            try:
                img = OpenpyxlImage(logo_path)
                img.width = 160
                img.height = 32
                size = XDRPositiveSize2D(pixels_to_EMU(img.width), pixels_to_EMU(img.height))
                right_col_idx = max(0, len(headers) - 2) if len(headers) >= 2 else 0
                marker = AnchorMarker(col=right_col_idx, colOff=pixels_to_EMU(15), row=0, rowOff=pixels_to_EMU(4))
                img.anchor = OneCellAnchor(_from=marker, ext=size)
                ws.add_image(img)
            except Exception as e:
                print(f"Warning: could not add logo to Excel generic report: {e}")

        header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9")
        )

        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=6, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        current_row = 7
        for row in rows:
            for col_idx, val in enumerate(row, 1):
                cell = ws.cell(row=current_row, column=col_idx, value=val)
                cell.border = thin_border
                if isinstance(val, (int, float)):
                    cell.number_format = '#,##0.00'
            current_row += 1

        for c in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(c)].width = 20

        # Print Setup for US Letter (8.5 x 11 in) Landscape Printing
        ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
        ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_margins.left = 0.4
        ws.page_margins.right = 0.4
        ws.page_margins.top = 0.5
        ws.page_margins.bottom = 0.5
        ws.print_options.horizontalCentered = True
        ws.print_title_rows = '6:6'
        ws.oddFooter.left.text = f"RED PUBLICIS — {report_title} (US Letter 8.5 x 11 in)"
        ws.oddFooter.right.text = "Page &P of &N"

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output
